# -*- coding: utf-8 -*-
"""Dokumen tunggal: seluruh soal yang lulus + soal tambahan baru, tersusun per mata pelajaran > materi > TP, dengan kunci dan pembahasan."""
import json, csv, collections, sys, re
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

bank = json.load(open('bank_final_lulus.json', encoding='utf-8'))
reg = json.load(open('tp_registry.json', encoding='utf-8'))
cov = list(csv.DictReader(open('cakupan_tp.csv', encoding='utf-8')))
summ = json.load(open('ringkasan.json', encoding='utf-8'))
ORDER = ['Fikih', 'Al-Quran Hadis', 'Bahasa Arab', 'Akidah Akhlak', 'Bahasa Indonesia', 'Matematika', 'Pancasila', 'PJOK', 'Bahasa Inggris']
NAMA = {'Al-Quran Hadis': "Al-Qur'an Hadis"}
TIPE = {'multiple_choice': 'pilihan ganda', 'multiple_choice_visual': 'pilihan ganda bergambar', 'contextual': 'pilihan ganda kontekstual', 'true_false': 'benar-salah',
        'fill_in_short': 'isian singkat', 'completion': 'melengkapi', 'matching': 'menjodohkan', 'ordering': 'mengurutkan', 'grouping': 'mengelompokkan', 'performance_task': 'tugas unjuk kerja', 'short_answer': 'isian singkat'}

doc = Document()
sec = doc.sections[0]; sec.page_width = Cm(21); sec.page_height = Cm(29.7)
for a in ('left_margin', 'right_margin'): setattr(sec, a, Cm(2))
sec.top_margin = Cm(1.8); sec.bottom_margin = Cm(1.8)
st = doc.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(10)
st.element.rPr.rFonts.set(qn('w:cs'), 'Arial'); st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
st.paragraph_format.space_after = Pt(1); st.paragraph_format.space_before = Pt(0)
for h, sz in (('Heading 1', 16), ('Heading 2', 13), ('Heading 3', 11)):
    s = doc.styles[h]; s.font.name = 'Calibri'; s.font.size = Pt(sz); s.font.color.rgb = RGBColor(0x1F, 0x3A, 0x5F)
    s.paragraph_format.space_before = Pt(10 if h != 'Heading 3' else 6); s.paragraph_format.space_after = Pt(3); s.paragraph_format.keep_with_next = True

def para(text='', bold=False, italic=False, size=None, indent=0, color=None, style=None, after=None):
    p = doc.add_paragraph(style=style)
    if indent: p.paragraph_format.left_indent = Cm(indent)
    if after is not None: p.paragraph_format.space_after = Pt(after)
    if text:
        r = p.add_run(text); r.bold = bold; r.italic = italic
        if size: r.font.size = Pt(size)
        if color: r.font.color.rgb = RGBColor(*color)
    return p

def runs(p, parts):
    for t, kw in parts:
        r = p.add_run(t); r.bold = kw.get('b', False); r.italic = kw.get('i', False)
        if kw.get('c'): r.font.color.rgb = RGBColor(*kw['c'])
        if kw.get('s'): r.font.size = Pt(kw['s'])

def table(rows, header, widths=None, size=8.5):
    t = doc.add_table(rows=1, cols=len(header)); t.style = 'Light Grid Accent 1'
    for i, h in enumerate(header):
        c = t.rows[0].cells[i]; c.text = ''; r = c.paragraphs[0].add_run(str(h)); r.bold = True; r.font.size = Pt(size)
    for row in rows:
        cs = t.add_row().cells
        for i, v in enumerate(row):
            cs[i].text = ''; r = cs[i].paragraphs[0].add_run(str(v)); r.font.size = Pt(size)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths): row.cells[i].width = Cm(w)
    return t

def toc():
    p = doc.add_paragraph(); r = p.add_run()
    for typ, txt in (('begin', None), (None, 'TOC \\o "1-2" \\h \\z \\u'), ('separate', None), (None, 'Klik kanan > Perbarui bidang untuk menampilkan daftar isi.'), ('end', None)):
        if typ:
            f = OxmlElement('w:fldChar'); f.set(qn('w:fldCharType'), typ); r._r.append(f)
        elif txt.startswith('TOC'):
            i = OxmlElement('w:instrText'); i.set(qn('xml:space'), 'preserve'); i.text = txt; r._r.append(i)
        else:
            t = OxmlElement('w:t'); t.text = txt; r._r.append(t)

def kunci_text(q):
    t, a, o = q['question_type'], q['answer'], q['options']
    if t in ('multiple_choice', 'multiple_choice_visual', 'contextual'):
        return next((x for x in o if x.startswith(a + '.')), a)
    if t == 'ordering': return ' → '.join(a)
    if t == 'matching': return '; '.join(f'{k} — {v}' for k, v in a.items())
    if t == 'grouping': return '; '.join(f'{k}: ' + ', '.join(v) for k, v in a.items())
    return str(a)

def soal(no, q):
    p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(5); p.paragraph_format.keep_with_next = True
    baru = q.get('sumber') != 'BANK_PASS'
    runs(p, [(f'{no}. ', {'b': True}), (f'[{q["id"]}] ', {'b': True, 'c': (0x1F, 0x3A, 0x5F), 's': 8.5}),
             (q['question'], {}), (f'  ({TIPE.get(q["question_type"], q["question_type"])}; {q["difficulty"]})', {'i': True, 'c': (0x77, 0x77, 0x77), 's': 8.5})]
         + ([('  [BARU]', {'b': True, 'c': (0xC0, 0x50, 0x00), 's': 8.5})] if baru else []))
    if q.get('visual_prompt'): para('Gambar: ' + q['visual_prompt'], italic=True, size=9, indent=0.6, color=(0x55, 0x55, 0x55))
    t, o = q['question_type'], q['options']
    if t in ('multiple_choice', 'multiple_choice_visual', 'contextual'):
        for x in o: para(x, indent=0.8, after=0)
    elif t == 'true_false': para('Benar / Salah', indent=0.8, after=0)
    elif t == 'ordering': para('Butir (acak): ' + ' | '.join(o), indent=0.8, after=0)
    elif t == 'matching': para('Kolom kiri: ' + ' | '.join(o['kolom_kiri']) + '   Kolom kanan: ' + ' | '.join(o['kolom_kanan']), indent=0.8, after=0)
    elif t == 'grouping': para('Kelompok: ' + ' | '.join(o['kelompok']), indent=0.8, after=0)
    p = doc.add_paragraph(); p.paragraph_format.left_indent = Cm(0.8)
    runs(p, [('Kunci: ', {'b': True, 'c': (0x1B, 0x6E, 0x3C), 's': 9}), (kunci_text(q), {'s': 9}), ('   Pembahasan: ', {'b': True, 'c': (0x55, 0x55, 0x55), 's': 9}), (q.get('explanation') or '', {'s': 9})])

# ---------------------------------------------------------------- halaman depan
p = para('BANK SOAL KELAS 1 MI DARUL ARIFIN', bold=True, size=22, color=(0x1F, 0x3A, 0x5F)); p.alignment = 1
p = para('Fase A — Seluruh soal yang lulus dan cakupan seluruh Tujuan Pembelajaran (TP)', size=13); p.alignment = 1
p = para('Disusun 15 Oktober 2026. Dokumen tunggal berisi soal, kunci jawaban, dan pembahasan.', italic=True, size=10); p.alignment = 1
doc.add_paragraph()
doc.add_heading('Ringkasan', 1)
n_baru = summ['soal_baru_lolos']
for s in (f'Total soal dalam dokumen ini: {summ["total_soal_final"]:,}'.replace(',', '.'),
          f'Soal lama berstatus PASS pada audit kurikulum: {summ["soal_lama_pass"]:,}'.replace(',', '.'),
          f'Soal baru yang ditambahkan agar semua TP tercakup dan merata: {n_baru} (ditandai [BARU])',
          f'Jumlah TP: {summ["tp_total"]} (248 TP yang sudah ada di bank + 4 TP ATP Bahasa Indonesia yang sebelumnya belum punya soal)',
          f'TP tanpa soal: {len(summ["tp_tanpa_soal"])}; TP di bawah target pemerataan: {len(summ["tp_di_bawah_target"])}'):
    para('• ' + s, indent=0.4)
doc.add_heading('Sumber acuan dan cara pemerataan', 2)
for s in ('Acuan utama adalah dokumen CP, TP, dan ATP Kelas 1 (Fase A) MI Darul Arifin, yaitu berkas ATP per mata pelajaran di repositori. KKTP tidak dipakai sebagai penentu cakupan; 330 soal tambahan Bahasa Arab berdasarkan KKTP (Hobiku, Buah, Warna) TIDAK dimasukkan karena bab itu tidak ada pada ATP.',
          'Target per TP = median jumlah soal PASS per TP pada mata pelajaran yang sama (dibulatkan ke atas). TP yang berada di bawah target ditambah soal baru sampai mencapai target; TP yang sudah di atas target tidak dikurangi.',
          'Soal lama yang berstatus NEEDS_REVIEW (1.302 soal) tidak dimasukkan karena menunggu keputusan guru/ahli. Soal lama dipakai dalam versi setelah revisi audit (audit/FINAL).',
          'Soal baru dibuat oleh model AI dari rumusan TP pada ATP, divalidasi otomatis (bentuk, pengecoh, duplikasi terhadap seluruh bank), diselesaikan secara buta oleh penyelesai AI lain, dan yang berbeda atau bermasalah diperiksa ulang oleh adjudikator AI. Soal yang tidak lolos tidak dimasukkan.'):
    para('• ' + s, indent=0.4)
doc.add_heading('Keterbatasan', 2)
for s in ('Semua pemeriksaan dilakukan oleh model AI; tidak ada persetujuan guru atau ahli. Label "lulus" berarti lolos pemeriksaan otomatis dan tinjau baca AI, bukan pengesahan.',
          'Materi agama (Fikih, Akidah Akhlak, Al-Qur\'an Hadis) dan Bahasa Arab, termasuk soal baru, perlu ditinjau ahli sebelum dipakai untuk penilaian resmi.',
          'Kecukupan "tercakup" diukur dari jumlah soal yang dipetakan pada setiap TP, bukan bukti bahwa soal tulis dapat mengukur TP yang menuntut praktik, lisan, hafalan, atau sikap. Untuk TP seperti itu tetap diperlukan rubrik observasi dan tugas unjuk kerja guru.',
          'Teks TP pada bank memuat frasa tambahan (mis. "sebagai wujud cinta ...") pada sebagian TP yang tidak ada di ATP; rumusan di dokumen ini mengikuti bank, kecuali 4 TP baru yang mengikuti ATP.'):
    para('• ' + s, indent=0.4)
doc.add_heading('Cakupan soal per mata pelajaran', 2)
rows = []
for s in ORDER:
    v = summ['per_mapel'][s]
    rows.append([NAMA.get(s, s), v['tp'], v['soal_final'], f'{v["awal"]["min"]} / {v["awal"]["median"]:g} / {v["awal"]["maks"]}', f'{v["akhir"]["min"]} / {v["akhir"]["median"]:g} / {v["akhir"]["maks"]}'])
table(rows, ['Mata pelajaran', 'Jumlah TP', 'Soal final', 'Soal per TP sebelum (min/median/maks)', 'Soal per TP sesudah (min/median/maks)'], [3.6, 2, 2.2, 4.6, 4.6], 9)
doc.add_paragraph()
doc.add_heading('Daftar isi', 2)
para('Gunakan panel navigasi (Heading 1-3) untuk berpindah antar bagian. Isi dokumen: ', size=9)
_by = collections.defaultdict(list)
for q in bank: _by[q['tp_id']].append(q)
for s_ in ORDER:
    mats = collections.OrderedDict()
    for e in reg:
        if e['subject'] == s_: mats[e['material']] = mats.get(e['material'], 0) + len(_by[e['tp_id']])
    para('• ' + NAMA.get(s_, s_) + ': ' + '; '.join(f'{m} ({n})' for m, n in mats.items()), indent=0.4, size=8.5)
para('• Lampiran: Cakupan soal pada setiap TP (252 TP), di akhir dokumen.', indent=0.4, size=8.5)

# ---------------------------------------------------------------- isi
by = collections.defaultdict(list)
for q in bank: by[q['tp_id']].append(q)
tp_order = {e['tp_id']: i for i, e in enumerate(reg)}
covd = {c['tp_id']: c for c in cov}
for s in ORDER:
    doc.add_page_break()
    doc.add_heading(f'{NAMA.get(s, s)}', 1)
    tps = [e for e in reg if e['subject'] == s]
    jml = sum(len(by[e['tp_id']]) for e in tps)
    para(f'{len(tps)} TP, {jml} soal.', italic=True, size=9)
    cur = None; no = 0
    for e in tps:
        if e['material'] != cur:
            cur = e['material']; doc.add_heading(cur, 2)
        c = covd[e['tp_id']]
        doc.add_heading(f'{e["tp_id"]} — {e["tp"]}', 3)
        para(f'Soal: {len(by[e["tp_id"]])} (lama {c["pass_awal"]}, baru {c["tambahan_lolos"]}; target {c["target"]})' + ('  — TP dari ATP yang sebelumnya belum punya soal' if e['baru'] else ''), italic=True, size=8.5, color=(0x55, 0x55, 0x55))
        for i, q in enumerate(sorted(by[e['tp_id']], key=lambda x: x['id']), 1): soal(i, q)
doc.add_page_break()
doc.add_heading('Lampiran. Cakupan soal pada setiap TP', 1)
para('Target = median jumlah soal PASS per TP pada mata pelajaran yang sama. Kolom "lama" = soal PASS dari bank yang sudah ada; "baru" = soal tambahan yang lolos verifikasi.', italic=True, size=9)
table([[NAMA.get(c['subject'], c['subject']), c['tp_id'], c['pass_awal'], c['tambahan_lolos'], c['total_final'], c['target'], c['tercapai']] for c in cov],
      ['Mapel', 'TP', 'Lama', 'Baru', 'Total', 'Target', 'Tercapai'], [3.2, 3.6, 1.5, 1.5, 1.5, 1.6, 2], 8)
doc.core_properties.title = 'Bank Soal Kelas 1 MI Darul Arifin - Soal Lulus dan Cakupan TP'
doc.core_properties.author = 'Disusun dengan bantuan AI; menunggu tinjauan guru/ahli'
doc.save(sys.argv[1])
print('tersimpan', sys.argv[1])
